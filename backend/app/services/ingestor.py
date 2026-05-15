import io
import logging
import uuid
from typing import Optional

import pdfplumber
import tiktoken
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from openai import AsyncOpenAI, APIError, APITimeoutError
from pptx import Presentation
from sqlalchemy import update
from tenacity import retry, retry_if_exception_type, stop_after_attempt, wait_exponential

from app.core.config import settings
from app.core.database import AsyncSessionLocal, Document, DocumentChunk

_client = AsyncOpenAI(api_key=settings.openai_api_key)
_enc = tiktoken.get_encoding("cl100k_base")
logger = logging.getLogger(__name__)

_RETRY = dict(
    stop=stop_after_attempt(3),
    wait=wait_exponential(multiplier=1, min=2, max=8),
    retry=retry_if_exception_type((APIError, APITimeoutError)),
    reraise=True,
)

CHUNK_SIZE = 512
CHUNK_OVERLAP = 50
EMBED_BATCH_SIZE = 100


def _split_text(text: str) -> list[str]:
    tokens = _enc.encode(text)
    chunks: list[str] = []
    start = 0
    while start < len(tokens):
        end = min(start + CHUNK_SIZE, len(tokens))
        chunks.append(_enc.decode(tokens[start:end]))
        start += CHUNK_SIZE - CHUNK_OVERLAP
    return chunks


@retry(**_RETRY)
async def _embed_batch(texts: list[str]) -> list[list[float]]:
    resp = await _client.embeddings.create(
        model=settings.embedding_model,
        input=texts,
        dimensions=settings.embedding_dimensions,
    )
    return [item.embedding for item in resp.data]


def _extract_pdf(file_bytes: bytes) -> list[dict]:
    records = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            text = text.strip()
            if len(text) < 40:
                continue
            for chunk in _split_text(text):
                records.append({"text": chunk, "page": page_num, "section": "body"})
    return records


def _extract_docx(file_bytes: bytes) -> list[dict]:
    records = []
    doc = DocxDocument(io.BytesIO(file_bytes))
    current_section = "body"
    buffer = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style.name.startswith("Heading"):
            current_section = text
        buffer.append(text)

    full_text = "\n".join(buffer)
    for chunk in _split_text(full_text):
        records.append({"text": chunk, "page": None, "section": current_section})
    return records


def _extract_pptx(file_bytes: bytes) -> list[dict]:
    records = []
    prs = Presentation(io.BytesIO(file_bytes))
    for slide_num, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
        full_text = "\n".join(texts).strip()
        if len(full_text) < 40:
            continue
        for chunk in _split_text(full_text):
            records.append({"text": chunk, "page": slide_num, "section": "slide"})
    return records


def _extract_html(file_bytes: bytes) -> list[dict]:
    soup = BeautifulSoup(file_bytes, "lxml")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = soup.get_text(separator="\n").strip()
    records = []
    for chunk in _split_text(text):
        records.append({"text": chunk, "page": None, "section": "body"})
    return records


def _extract_text(file_bytes: bytes) -> list[dict]:
    text = file_bytes.decode("utf-8", errors="ignore").strip()
    records = []
    for chunk in _split_text(text):
        records.append({"text": chunk, "page": None, "section": "body"})
    return records


def _parse(file_bytes: bytes, filename: str) -> list[dict]:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        return _extract_pdf(file_bytes)
    if ext == "docx":
        return _extract_docx(file_bytes)
    if ext == "pptx":
        return _extract_pptx(file_bytes)
    if ext in ("html", "htm"):
        return _extract_html(file_bytes)
    return _extract_text(file_bytes)  # txt, md, fallback


async def process_document(
    file_bytes: bytes,
    filename: str,
    doc_id: str,
    dept: str,
    sensitivity: str,
    allowed_groups: list[str],
) -> None:
    async with AsyncSessionLocal() as db:
        try:
            records = _parse(file_bytes, filename)

            all_embeddings: list[list[float]] = []
            for i in range(0, len(records), EMBED_BATCH_SIZE):
                batch_texts = [r["text"] for r in records[i: i + EMBED_BATCH_SIZE]]
                all_embeddings.extend(await _embed_batch(batch_texts))

            for record, embedding in zip(records, all_embeddings):
                db.add(
                    DocumentChunk(
                        id=str(uuid.uuid4()),
                        doc_id=doc_id,
                        filename=filename,
                        page=record["page"],
                        section=record["section"],
                        content=record["text"],
                        embedding=embedding,
                        allowed_groups=allowed_groups,
                        dept=dept,
                        sensitivity=sensitivity,
                    )
                )

            await db.execute(
                update(Document)
                .where(Document.id == doc_id)
                .values(status="indexed", chunk_count=len(records))
            )
            await db.commit()

        except Exception as exc:
            logger.exception("ingestion_failed", extra={"doc_id": doc_id, "filename": filename})
            await db.execute(
                update(Document)
                .where(Document.id == doc_id)
                .values(status="error")
            )
            await db.commit()
            raise
